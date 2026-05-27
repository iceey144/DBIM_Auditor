# =========================================================
# DBIM ANNEXURE A.2 AUDITOR
# ULTRA REFINED ACCURACY VERSION
# =========================================================

import asyncio
import os
import re
import cv2

from playwright.async_api import async_playwright
from docx import Document
from docx.shared import Inches

# =========================================================
# FOLDERS
# =========================================================

FOLDERS = [

    "screenshots",
    "reports"

]

for folder in FOLDERS:

    os.makedirs(folder, exist_ok=True)

# =========================================================
# PERSONA DEFINITIONS
# =========================================================

PERSONA_KEYWORDS = {

    "Students": [

        "student",
        "students",
        "scholarship",
        "education",
        "college",
        "university"

    ],

    "Farmers": [

        "farmer",
        "farmers",
        "agriculture",
        "kisan",
        "crop"

    ],

    "Citizens": [

        "citizen",
        "citizens",
        "resident",
        "public service"

    ],

    "Healthcare": [

        "doctor",
        "patient",
        "healthcare",
        "medical"

    ],

    "Businesses": [

        "business",
        "startup",
        "industry",
        "enterprise"

    ]

}

# =========================================================
# NORMALIZE
# =========================================================

def normalize(text):

    return re.sub(

        r"\s+",
        " ",
        str(text).lower()

    ).strip()

# =========================================================
# SCREENSHOT
# =========================================================

async def take_screenshot(

    page,
    path

):

    await page.screenshot(

        path=path,
        full_page=True

    )

# =========================================================
# VALID BOX
# =========================================================

def valid_box(box):

    if not box:
        return False

    w = box["width"]
    h = box["height"]

    # tiny noise
    if w < 55:
        return False

    if h < 22:
        return False

    # giant sections
    if w > 450:
        return False

    if h > 180:
        return False

    return True

# =========================================================
# DRAW BOXES
# =========================================================

def draw_detection_boxes(

    image_path,
    detections,
    output_path

):

    img = cv2.imread(image_path)

    if img is None:
        return

    for d in detections:

        x = int(d["x"])
        y = int(d["y"])

        w = int(d["w"])
        h = int(d["h"])

        label = d["label"]

        color = d["color"]

        cv2.rectangle(

            img,

            (x,y),
            (x+w,y+h),

            color,
            3

        )

        cv2.putText(

            img,

            label,

            (x,y-10),

            cv2.FONT_HERSHEY_SIMPLEX,
            0.65,
            color,
            2

        )

    cv2.imwrite(output_path, img)

# =========================================================
# FILTER ELEMENTS
# =========================================================

async def should_ignore(el, box):

    try:

        if not await el.is_visible():
            return True

        if not valid_box(box):
            return True

        html = await el.evaluate(
            "(e) => e.outerHTML"
        )

        lower_html = normalize(html)

        # ignore footer spam
        if "footer" in lower_html:
            return True

        # ignore hidden elements
        style = await el.get_attribute("style")

        if style:

            style = normalize(style)

            if "display:none" in style:
                return True

            if "visibility:hidden" in style:
                return True

        return False

    except:
        return True

# =========================================================
# DETECT PERSONA NAVIGATION
# =========================================================

async def detect_persona_navigation(page):

    detections = []

    evidence = []

    seen = set()

    selectors = [

        "nav a",
        "header a",
        ".menu a",
        ".navbar a",
        ".nav-link",
        "button"

    ]

    for selector in selectors:

        try:

            elements = await page.query_selector_all(
                selector
            )

            for el in elements:

                try:

                    text = await el.inner_text()

                    if not text:
                        continue

                    lower = normalize(text)

                    # avoid paragraphs
                    if len(lower.split()) > 6:
                        continue

                    matched = None

                    for persona, keywords in PERSONA_KEYWORDS.items():

                        for kw in keywords:

                            if kw in lower:

                                matched = persona
                                break

                        if matched:
                            break

                    if not matched:
                        continue

                    box = await el.bounding_box()

                    if await should_ignore(el, box):
                        continue

                    key = (

                        matched,
                        int(box["x"]),
                        int(box["y"])

                    )

                    if key in seen:
                        continue

                    seen.add(key)

                    evidence.append({

                        "persona": matched,
                        "text": text.strip()

                    })

                    detections.append({

                        "x": box["x"],
                        "y": box["y"],

                        "w": box["width"],
                        "h": box["height"],

                        "label": matched.upper(),

                        "color": (0,255,0)

                    })

                except:
                    pass

        except:
            pass

    return evidence, detections

# =========================================================
# DETECT PERSONA CARDS
# =========================================================

async def detect_persona_cards(page):

    detections = []

    evidence = []

    seen = set()

    selectors = [

        ".card",
        ".service-card",
        ".feature-card",
        ".tile",
        ".user-card",

        "[class*='card']",
        "[class*='service']",

        "section"

    ]

    for selector in selectors:

        try:

            elements = await page.query_selector_all(
                selector
            )

            for el in elements:

                try:

                    text = await el.inner_text()

                    if not text:
                        continue

                    lower = normalize(text)

                    # avoid large paragraphs
                    if len(lower.split()) > 16:
                        continue

                    matched = None

                    for persona, keywords in PERSONA_KEYWORDS.items():

                        for kw in keywords:

                            if kw in lower:

                                matched = persona
                                break

                        if matched:
                            break

                    if not matched:
                        continue

                    box = await el.bounding_box()

                    if await should_ignore(el, box):
                        continue

                    key = (

                        matched,
                        int(box["x"]),
                        int(box["y"])

                    )

                    if key in seen:
                        continue

                    seen.add(key)

                    evidence.append({

                        "persona": matched,
                        "text": text.strip()

                    })

                    detections.append({

                        "x": box["x"],
                        "y": box["y"],

                        "w": box["width"],
                        "h": box["height"],

                        "label": matched.upper(),

                        "color": (255,0,0)

                    })

                except:
                    pass

        except:
            pass

    return evidence, detections

# =========================================================
# DETECT PERSONA IMAGES
# =========================================================

async def detect_persona_images(page):

    detections = []

    evidence = []

    seen = set()

    keywords = [

        "student",
        "farmer",
        "doctor",
        "patient",
        "business",
        "citizen"

    ]

    try:

        images = await page.query_selector_all("img")

        for img in images:

            try:

                alt = await img.get_attribute("alt")

                src = await img.get_attribute("src")

                combined = normalize(

                    str(alt) + " " + str(src)

                )

                matched = None

                for kw in keywords:

                    if kw in combined:

                        matched = kw
                        break

                if not matched:
                    continue

                box = await img.bounding_box()

                if await should_ignore(img, box):
                    continue

                key = (

                    matched,
                    int(box["x"]),
                    int(box["y"])

                )

                if key in seen:
                    continue

                seen.add(key)

                evidence.append({

                    "image": matched
                })

                detections.append({

                    "x": box["x"],
                    "y": box["y"],

                    "w": box["width"],
                    "h": box["height"],

                    "label": "PERSONA IMG",

                    "color": (255,120,0)

                })

            except:
                pass

    except:
        pass

    return evidence, detections

# =========================================================
# DOCX
# =========================================================

def generate_docx(results):

    doc = Document()

    doc.add_heading(

        "DBIM Annexure A.2 Audit Report",

        level=1

    )

    doc.add_paragraph(

        "User persona and audience-"
        "oriented design compliance audit."

    )

    for item in results:

        doc.add_heading(

            item["title"],
            level=2

        )

        para = doc.add_paragraph()

        para.add_run(

            f"STATUS: {item['status']}\n\n"

        ).bold = True

        para.add_run(item["reason"])

        try:

            doc.add_picture(

                item["image"],
                width=Inches(6)

            )

        except:
            pass

        doc.add_page_break()

    output = "reports/annexure_a2_report.docx"

    doc.save(output)

    print("\nDOCX GENERATED")
    print(output)

# =========================================================
# MAIN ENGINE
# =========================================================

async def engine():

    url = input(
        "\nEnter Website URL: "
    )

    results = []

    async with async_playwright() as p:

        browser = await p.chromium.launch(

            headless=False,

            slow_mo=400,

            args=[

                "--disable-blink-features=AutomationControlled",

                "--disable-dev-shm-usage",

                "--no-sandbox",

                "--disable-gpu"

            ]

        )

        context = await browser.new_context(

            viewport={

                "width":1280,
                "height":720

            },

            ignore_https_errors=True

        )

        page = await context.new_page()

        print("\nOPENING WEBSITE...")

        try:

            await page.goto(

                url,

                wait_until="domcontentloaded",

                timeout=120000

            )

            await page.wait_for_load_state(
                "networkidle"
            )

        except Exception as e:

            print(f"\nLOAD FAILED:\n{e}")

            await browser.close()

            return

        # =================================================
        # HUMAN-LIKE SCROLLING
        # =================================================

        await page.mouse.wheel(0, 2500)
        await page.wait_for_timeout(2000)

        await page.mouse.wheel(0, 2500)
        await page.wait_for_timeout(2000)

        await page.mouse.wheel(0, -2500)
        await page.wait_for_timeout(2000)

        print("\nWEBSITE LOADED")

        screenshot = "screenshots/full_page.png"

        await take_screenshot(

            page,
            screenshot

        )

        # =================================================
        # PERSONA NAVIGATION
        # =================================================

        print("\nTESTING PERSONA NAVIGATION")

        nav_evidence, nav_boxes = await detect_persona_navigation(page)

        nav_compliant = len(nav_evidence) > 0

        nav_image = screenshot

        if nav_compliant:

            nav_image = "screenshots/nav_detection.png"

            draw_detection_boxes(

                screenshot,
                nav_boxes,
                nav_image

            )

            lines = []

            for e in nav_evidence[:5]:

                lines.append(

                    f"- Navigation item detected: '{e['text']}'"
                )

            reason = f"""
RULE: Persona Navigation

STATUS: COMPLIANT

Reason:
Role-oriented navigation structures
targeting specific audience groups
were detected.

Evidence:
{chr(10).join(lines)}
"""

        else:

            reason = """
RULE: Persona Navigation

STATUS: NON-COMPLIANT

Reason:
No role-oriented navigation
structures targeting specific
user personas were detected.
"""

        results.append({

            "title":
            "Persona Navigation",

            "status":

            "COMPLIANT"

            if nav_compliant

            else "NON-COMPLIANT",

            "reason":
            reason,

            "image":
            nav_image

        })

        # =================================================
        # PERSONA CARDS
        # =================================================

        print("\nTESTING PERSONA CARDS")

        card_evidence, card_boxes = await detect_persona_cards(page)

        card_compliant = len(card_evidence) > 0

        card_image = screenshot

        if card_compliant:

            card_image = "screenshots/card_detection.png"

            draw_detection_boxes(

                screenshot,
                card_boxes,
                card_image

            )

            lines = []

            for e in card_evidence[:5]:

                lines.append(

                    f"- Persona section detected: '{e['text'][:70]}'"
                )

            reason = f"""
RULE: Persona-Oriented Sections

STATUS: COMPLIANT

Reason:
Persona-oriented service/content
sections targeting user groups
were detected.

Evidence:
{chr(10).join(lines)}
"""

        else:

            reason = """
RULE: Persona-Oriented Sections

STATUS: NON-COMPLIANT

Reason:
No persona-oriented service/content
sections were detected.
"""

        results.append({

            "title":
            "Persona-Oriented Sections",

            "status":

            "COMPLIANT"

            if card_compliant

            else "NON-COMPLIANT",

            "reason":
            reason,

            "image":
            card_image

        })

        # =================================================
        # PERSONA IMAGERY
        # =================================================

        print("\nTESTING PERSONA IMAGERY")

        image_evidence, image_boxes = await detect_persona_images(page)

        image_compliant = len(image_evidence) > 0

        image_output = screenshot

        if image_compliant:

            image_output = "screenshots/image_detection.png"

            draw_detection_boxes(

                screenshot,
                image_boxes,
                image_output

            )

            lines = []

            for e in image_evidence[:5]:

                lines.append(

                    f"- Persona image detected: '{e['image']}'"
                )

            reason = f"""
RULE: Persona Imagery

STATUS: COMPLIANT

Reason:
Persona-oriented imagery related
to specific audience groups
was detected.

Evidence:
{chr(10).join(lines)}
"""

        else:

            reason = """
RULE: Persona Imagery

STATUS: NON-COMPLIANT

Reason:
No persona-oriented imagery
was detected.
"""

        results.append({

            "title":
            "Persona Imagery",

            "status":

            "COMPLIANT"

            if image_compliant

            else "NON-COMPLIANT",

            "reason":
            reason,

            "image":
            image_output

        })

        # =================================================
        # GENERATE DOCX
        # =================================================

        generate_docx(results)

        await browser.close()

# =========================================================
# RUN
# =========================================================

asyncio.run(
    engine()
)